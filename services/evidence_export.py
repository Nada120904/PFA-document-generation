"""Structured, source-attributed exports shared by demo and production."""

from __future__ import annotations

import csv
import re
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape


@dataclass(frozen=True)
class Evidence:
    source: str
    excerpt: str


def evidence_from_lines(lines: list[str]) -> list[Evidence]:
    entries: list[Evidence] = []
    for line in lines:
        source, marker, excerpt = line.partition(": ")
        if marker and excerpt.strip():
            entries.append(Evidence(Path(source.strip()).name, excerpt.strip()))
    return entries


def grouped(entries: list[Evidence]) -> OrderedDict[str, list[str]]:
    groups: OrderedDict[str, list[str]] = OrderedDict()
    for entry in entries:
        if entry.excerpt and entry.excerpt not in groups.setdefault(entry.source, []):
            groups[entry.source].append(entry.excerpt)
    return groups


def source_label(source: str) -> str:
    return re.sub(r"[_-]+", " ", Path(source).stem).strip().title()


def _highlights(entries: list[Evidence]) -> list[Evidence]:
    """Pick useful first facts across sources without inventing a summary."""
    result: list[Evidence] = []
    seen: set[str] = set()
    for entry in entries:
        if entry.source not in seen:
            result.append(entry)
            seen.add(entry.source)
        if len(result) == 4:
            break
    if len(result) < 3:
        for entry in entries:
            if entry not in result:
                result.append(entry)
            if len(result) == 3:
                break
    return result


def write_txt(path: Path, title: str, entries: list[Evidence]) -> None:
    groups = grouped(entries)
    count = sum(map(len, groups.values()))
    rule = "=" * min(max(len(title), 28), 78)
    lines = [title, rule, "", f"SOURCE FILES  {len(groups)}", f"EVIDENCE ITEMS  {count}", ""]
    lines += ["HIGHLIGHTS", "----------"]
    for number, item in enumerate(_highlights(entries), 1):
        lines += [f"{number:02d}. {item.excerpt}", f"    Source: {item.source}", ""]
    lines += ["SOURCE INDEX", "------------"]
    for number, (source, excerpts) in enumerate(groups.items(), 1):
        lines.append(f"{number:02d}. {source}  |  {len(excerpts)} evidence items")
    lines += ["", "EVIDENCE BY DOCUMENT", "--------------------"]
    for number, (source, excerpts) in enumerate(groups.items(), 1):
        lines += ["", f"{number:02d}  {source_label(source).upper()}", f"    {source}"]
        for item_number, excerpt in enumerate(excerpts, 1):
            lines.append(f"    {item_number:02d}. {excerpt}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_csv(path: Path, entries: list[Evidence]) -> list[list[str]]:
    groups = grouped(entries)
    rows: list[list[str]] = []
    number = 1
    for source, excerpts in groups.items():
        for excerpt in excerpts:
            rows.append([f"E{number:03d}", source_label(source), source, Path(source).suffix.lstrip(".").upper(), excerpt])
            number += 1
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["ID", "Section", "Source file", "File type", "Evidence excerpt"])
        writer.writerows(rows)
    return rows


def write_docx(path: Path, title: str, entries: list[Evidence]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    groups = grouped(entries)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(37, 53, 62)
    normal.paragraph_format.space_after = Pt(4)
    for name, size, before, after in (("Title", 23, 0, 7), ("Heading 1", 14, 16, 7), ("Heading 2", 11, 11, 4)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0) if name == "Title" else RGBColor(0, 105, 111)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if name == "Title":
            paragraph_properties = style.element.get_or_add_pPr()
            borders = paragraph_properties.find(qn("w:pBdr"))
            if borders is None:
                borders = OxmlElement("w:pBdr")
                paragraph_properties.append(borders)
            bottom = borders.find(qn("w:bottom"))
            if bottom is None:
                bottom = OxmlElement("w:bottom")
                borders.append(bottom)
            bottom.set(qn("w:val"), "nil")
        elif name == "Heading 2":
            style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "PFA 2026  /  DOCUMENT INTELLIGENCE"
    header.style = normal
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor(0, 105, 111)
    footer = section.footer.paragraphs[0]
    footer.text = "Source evidence report  |  Generated from uploaded documents"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 116, 124)

    document.add_paragraph(title, style="Title")
    lead = document.add_paragraph()
    lead.add_run(f"{len(groups)} source files  ·  {sum(map(len, groups.values()))} evidence items").bold = True
    lead.runs[0].font.color.rgb = RGBColor(0, 105, 111)
    document.add_paragraph("A structured review of the passages contained in the uploaded files.")

    document.add_heading("Highlights", level=1)
    for item in _highlights(entries):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item.excerpt)
        source = paragraph.add_run(f"  [{item.source}]")
        source.italic = True
        source.font.color.rgb = RGBColor(91, 111, 119)

    document.add_heading("Source index", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ("Document", "Type", "Items")):
        cell.text = value
    for source, excerpts in groups.items():
        cells = table.add_row().cells
        for cell, value in zip(cells, (source, Path(source).suffix.lstrip(".").upper(), str(len(excerpts)))):
            cell.text = value

    document.add_heading("Evidence by document", level=1)
    for group_number, (source, excerpts) in enumerate(groups.items()):
        if group_number == 1 and len(groups) >= 4 and sum(map(len, groups.values())) >= 30:
            document.add_page_break()
        document.add_heading(source_label(source), level=2)
        source_paragraph = document.add_paragraph(source)
        source_paragraph.paragraph_format.keep_with_next = True
        for number, excerpt in enumerate(excerpts, 1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            if number == 1:
                paragraph.paragraph_format.keep_with_next = True
            marker = paragraph.add_run(f"{number:02d}  ")
            marker.bold = True
            marker.font.color.rgb = RGBColor(0, 105, 111)
            paragraph.add_run(excerpt)
    document.save(path)


def write_pdf(path: Path, title: str, entries: list[Evidence]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    groups = grouped(entries)
    navy = colors.HexColor("#15333e")
    teal = colors.HexColor("#00696f")
    muted = colors.HexColor("#5b707a")
    light = colors.HexColor("#ecf5f5")
    styles = {
        "eyebrow": ParagraphStyle("Eyebrow", fontName="Helvetica-Bold", fontSize=8, leading=11, textColor=teal, spaceAfter=9),
        "title": ParagraphStyle("ReportTitle", fontName="Helvetica-Bold", fontSize=22, leading=27, textColor=navy, spaceAfter=8),
        "lead": ParagraphStyle("Lead", fontName="Helvetica", fontSize=10, leading=15, textColor=muted, spaceAfter=14),
        "h1": ParagraphStyle("Section", fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=teal, spaceBefore=18, spaceAfter=8),
        "h2": ParagraphStyle("Source", fontName="Helvetica-Bold", fontSize=10.5, leading=14, textColor=navy, spaceBefore=14, spaceAfter=3),
        "body": ParagraphStyle("Evidence", fontName="Helvetica", fontSize=9.3, leading=13.6, textColor=navy, spaceAfter=7),
        "small": ParagraphStyle("Small", fontName="Helvetica", fontSize=8, leading=11, textColor=muted),
    }
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=19*mm, bottomMargin=18*mm, title=title,
                            author="PFA 2026 Document Intelligence")
    story: list = [
        Paragraph("PFA 2026  /  SOURCE EVIDENCE REPORT", styles["eyebrow"]),
        Paragraph(escape(title), styles["title"]),
        Paragraph("A structured review of passages contained in the uploaded files.", styles["lead"]),
    ]
    metrics = Table([[
        Paragraph(f"<b>{len(groups)}</b><br/>SOURCE FILES", styles["body"]),
        Paragraph(f"<b>{sum(map(len, groups.values()))}</b><br/>EVIDENCE ITEMS", styles["body"]),
        Paragraph(f"<b>{len(_highlights(entries))}</b><br/>HIGHLIGHTS", styles["body"]),
    ]], colWidths=[56*mm, 56*mm, 56*mm])
    metrics.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), light), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 12), ("TOPPADDING", (0,0), (-1,-1), 10),
        ("BOTTOMPADDING", (0,0), (-1,-1), 8),
    ]))
    story += [metrics, Paragraph("Highlights", styles["h1"])]
    for number, item in enumerate(_highlights(entries), 1):
        story.append(Paragraph(f"<font color='#00696f'><b>{number:02d}</b></font>  {escape(item.excerpt)}<br/><font color='#5b707a' size='8'>{escape(item.source)}</font>", styles["body"]))
    story.append(Paragraph("Source index", styles["h1"]))
    rows = [[Paragraph("DOCUMENT", styles["small"]), Paragraph("TYPE", styles["small"]), Paragraph("ITEMS", styles["small"])]]
    for source, excerpts in groups.items():
        rows.append([Paragraph(escape(source), styles["body"]), Paragraph(Path(source).suffix.lstrip(".").upper(), styles["body"]), Paragraph(str(len(excerpts)), styles["body"])])
    index = Table(rows, colWidths=[110*mm, 29*mm, 29*mm], repeatRows=1)
    index.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), light), ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LINEBELOW", (0,0), (-1,-1), 0.3, colors.HexColor("#d7e5e7")),
        ("LEFTPADDING", (0,0), (-1,-1), 8), ("TOPPADDING", (0,0), (-1,-1), 6),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    story += [index, Paragraph("Evidence by document", styles["h1"])]
    for source, excerpts in groups.items():
        story.append(KeepTogether([
            Paragraph(escape(source_label(source)), styles["h2"]),
            Paragraph(escape(source), styles["small"]),
            Spacer(1, 3*mm),
        ]))
        for number, excerpt in enumerate(excerpts, 1):
            story.append(Paragraph(f"<font color='#00696f'><b>{number:02d}</b></font>  {escape(excerpt)}", styles["body"]))

    def page_chrome(canvas, document):
        canvas.saveState()
        width, height = A4
        canvas.setStrokeColor(colors.HexColor("#cfe0e3"))
        canvas.line(20*mm, height-14*mm, width-20*mm, height-14*mm)
        canvas.line(20*mm, 14*mm, width-20*mm, 14*mm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawString(20*mm, 10*mm, "Source evidence report")
        canvas.drawRightString(width-20*mm, 10*mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_chrome, onLaterPages=page_chrome)
