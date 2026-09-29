"""Local text extraction and small generated files for demo mode."""

from __future__ import annotations

import csv
import io
import re
from html.parser import HTMLParser
from pathlib import Path
from uuid import uuid4

from services.demo_service.db import GENERATED_DIR
from services.evidence_export import evidence_from_lines, write_csv, write_docx, write_pdf, write_txt


MAX_EXTRACTED_CHARS = 250_000


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


def safe_filename(filename: str) -> str:
    cleaned = Path(filename.replace("\\", "/")).name.strip()
    return cleaned or "document"


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".json"}:
        result = decode_text(data)
    elif suffix == ".csv":
        reader = csv.DictReader(io.StringIO(decode_text(data)))
        result = "\n".join(
            "; ".join(f"{key}: {value}" for key, value in row.items() if key and value)
            for row in reader if str(row.get("Metric", "")).casefold() != "demo run id"
        )
    elif suffix in {".html", ".htm"}:
        parser = _TextHTMLParser()
        parser.feed(decode_text(data))
        result = "\n".join(parser.parts)
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        result = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        result = "\n".join(chunks)
    else:
        raise ValueError(f"Demo text extraction does not support {suffix or 'this file type'}.")
    result = re.sub(r"[ \t]+", " ", result).strip()
    if not result:
        raise ValueError("No extractable text found. Image-only PDFs need OCR outside demo mode.")
    return result[:MAX_EXTRACTED_CHARS]


def make_generated_file(format_name: str, title: str, source_lines: list[str]) -> str:
    """Create a structured report based only on indexed document excerpts."""
    filename = f"{uuid4().hex}.{format_name}"
    target = GENERATED_DIR / filename
    entries = evidence_from_lines(source_lines)
    if not entries:
        raise ValueError("No indexed source excerpts were available for the export.")
    if format_name == "txt":
        write_txt(target, title, entries)
    elif format_name == "csv":
        write_csv(target, entries)
    elif format_name == "docx":
        write_docx(target, title, entries)
    elif format_name == "pdf":
        write_pdf(target, title, entries)
    else:
        raise ValueError("Unsupported generated format.")
    return filename
