"""Build four small, distinct upload examples for the local demo."""

from __future__ import annotations

import argparse
import uuid
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def build_samples(output_dir: Path, run_id: str) -> list[Path]:
    from docx import Document
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    output_dir.mkdir(parents=True, exist_ok=True)
    brief = output_dir / "project_brief.txt"
    brief.write_text(
        (ROOT / "demo_assets" / "project_brief.txt").read_text(encoding="utf-8")
        + f"\nDemo run ID: {run_id}\n", encoding="utf-8",
    )

    metrics = output_dir / "pilot_metrics.csv"
    metrics.write_text(
        (ROOT / "demo_assets" / "pilot_metrics.csv").read_text(encoding="utf-8")
        + f"Demo run ID,{run_id},Local recording\n", encoding="utf-8",
    )

    word = output_dir / "review_guide.docx"
    document = Document()
    document.add_heading("Pilot Review Guide", 0)
    for sentence in (
        "The quality reviewer is Salma Idrissi.",
        "The review owner is Amal Benali.",
        "The reviewer checks source excerpts in the preview before downloading an export.",
        "The reviewer compares every generated claim with the named uploaded file.",
        "The reviewer confirms that a requested subject appears in the indexed files before approving an export.",
        "A request without matching source material must return a no-data message and no file.",
        "The review covers PDF, Word, CSV, and TXT exports.",
        "The review checklist includes factual accuracy, readable structure, source names, and complete download links.",
        "The owner records any unsupported or unclear statement for correction before sharing a file.",
        "The target review turnaround is two business days.",
    ):
        document.add_paragraph(sentence)
    document.add_paragraph(f"Demo run ID: {run_id}")
    document.save(word)

    pdf = output_dir / "quality_criteria.pdf"
    page = canvas.Canvas(str(pdf), pagesize=letter)
    page.setFont("Helvetica-Bold", 17)
    page.drawString(50, 740, "Pilot Quality Criteria")
    page.setFont("Helvetica", 11)
    criteria = (
        "The pilot success criterion is zero unsupported claims in generated files.",
        "Every exported excerpt must be traceable to an indexed upload.",
        "Machine learning (ML) helps classify the pilot's uploaded documents.",
        "The ML workflow extracts text and checks source passages before export.",
        "ML classification organizes documents so reviewers can find relevant evidence.",
        "ML output is reviewed against the extracted text from the original files.",
        "ML-assisted retrieval uses indexed passages to answer document questions.",
        "ML does not authorize a claim when the indexed files contain no matching fact.",
        "The ML review checks the source name, passage wording, and requested subject.",
        "ML classification uses labels supported by the indexed document text.",
        "ML-assisted answers include source passages for reviewer verification.",
        "A missing subject must produce a no-data response without an export.",
        "A readable preview is required before a generated file is downloaded.",
        "The quality reviewer verifies every PDF, Word, CSV, and TXT output.",
    )
    y = 700
    for sentence in criteria:
        page.drawString(50, y, sentence)
        y -= 22
    page.drawString(50, y - 6, f"Demo run ID: {run_id}")
    page.save()
    return [brief, word, metrics, pdf]


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "demo_assets" / "sample_uploads")
    parser.add_argument("--run-id", default=uuid.uuid4().hex[:8])
    arguments = parser.parse_args()
    for path in build_samples(arguments.output_dir, arguments.run_id):
        print(path.resolve())
